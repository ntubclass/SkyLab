"""Rubric parser service - Parse text and office documents to Markdown."""

from __future__ import annotations

import contextlib
import io
import os
import shutil
import signal
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any

# Office 檔案就是 zip：解壓後的總大小與壓縮比都要先看過，
# 否則一份幾百 KB 的檔案可以在解析階段吃掉數 GB 記憶體。
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
# 單份 PDF 的頁數上限：pdfplumber 逐頁抽字與表格，頁數沒有上限時
# 一份檔案就能把 worker 佔住幾十分鐘。
MAX_PDF_PAGES = 200
# .doc 轉檔（LibreOffice）的逾時秒數
DOC_CONVERT_TIMEOUT_SECONDS = 30


def _guard_zip_document(file_bytes: bytes) -> None:
    """解壓炸彈防護：先讀 zip 目錄的宣告大小，再決定要不要解析。"""
    from app.core.i18n import t

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            total = sum(int(info.file_size) for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError(t("rubric_parser.corrupted_document")) from exc
    if total > MAX_ZIP_UNCOMPRESSED_BYTES or (
        file_bytes and total / len(file_bytes) > MAX_ZIP_COMPRESSION_RATIO
    ):
        raise ValueError(
            t(
                "rubric_parser.document_too_large",
                limit=MAX_ZIP_UNCOMPRESSED_BYTES // (1024 * 1024),
            )
        )


def parse_docx(file_bytes: bytes) -> str:
    """用 python-docx 解析 .docx，段落與表格都轉為 Markdown 文字。"""
    from docx import Document  # type: ignore

    _guard_zip_document(file_bytes)
    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []

    for block in _iter_block_items(doc):
        if block["type"] == "paragraph":
            text = block["text"].strip()
            if text:
                lines.append(text)
        elif block["type"] == "table":
            rows = [
                [_clean_cell(cell.text) for cell in row.cells]
                for row in block["table"].rows
            ]
            lines.extend(_markdown_table(rows))

    return "\n".join(lines)


def _clean_cell(value: object) -> str:
    """表格儲存格轉成單行文字：None 視為空字串，去頭尾空白、換行改空格。"""
    return str(value or "").strip().replace("\n", " ")


def _markdown_table(rows: list[list[str]]) -> list[str]:
    """把已清理過的儲存格轉成 Markdown 表格行；第一列當表頭，沒有列就回空。"""
    if not rows:
        return []
    header = rows[0]
    return [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
        *("| " + " | ".join(row) + " |" for row in rows[1:]),
    ]


def parse_pdf(file_bytes: bytes) -> str:
    """用 pdfplumber 解析純文字型 PDF，表格轉為 Markdown，其餘為純文字（排除表格區域）。"""
    from fastapi import HTTPException

    from app.core.i18n import t

    try:
        import pdfplumber  # type: ignore
    except ImportError as exc:
        raise HTTPException(
            status_code=503,
            detail=t("rubric_parser.pdfplumber_missing"),
        ) from exc

    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise ValueError(
                t("rubric_parser.pdf_too_many_pages", limit=MAX_PDF_PAGES)
            )
        for page in pdf.pages:
            # 先偵測此頁的表格與其 bbox，避免重複提取純文字
            tables = page.find_tables()
            table_bboxes: list[tuple[float, float, float, float]] = []
            for table in tables:
                data = table.extract()
                if not data:
                    continue
                table_bboxes.append(table.bbox)
                lines.extend(
                    _markdown_table(
                        [[_clean_cell(cell) for cell in row] for row in data]
                    )
                )

            # 純文字（排除表格區域）：只保留不在任何表格 bbox 內的文字
            words = page.extract_words() or []

            def _in_any_table_bbox(word: dict, table_bboxes: list = table_bboxes) -> bool:
                x0 = float(word.get("x0", 0.0))
                x1 = float(word.get("x1", 0.0))
                top = float(word.get("top", 0.0))
                bottom = float(word.get("bottom", 0.0))
                for tb_x0, tb_top, tb_x1, tb_bottom in table_bboxes:
                    if (
                        x0 >= tb_x0
                        and x1 <= tb_x1
                        and top >= tb_top
                        and bottom <= tb_bottom
                    ):
                        return True
                return False

            non_table_words = [w for w in words if not _in_any_table_bbox(w)]

            # 依行號或 y 座標重建純文字行
            lines_by_key: dict[int, list[dict]] = {}
            for w in non_table_words:
                key = w.get("line_number")
                if key is None:
                    # 退而求其次：用 top 四捨五入當作行的 key
                    key = int(round(float(w.get("top", 0.0))))
                key = int(key)
                lines_by_key.setdefault(key, []).append(w)

            page_text_lines: list[str] = []
            for key in sorted(lines_by_key.keys()):
                words_in_line = sorted(
                    lines_by_key[key], key=lambda ww: float(ww.get("x0", 0.0))
                )
                text_line = " ".join(
                    str(ww.get("text", "")).strip()
                    for ww in words_in_line
                    if str(ww.get("text", "")).strip()
                )
                if text_line:
                    page_text_lines.append(text_line)

            page_text = "\n".join(page_text_lines)
            if page_text.strip():
                lines.append(page_text.strip())
    return "\n".join(lines)


def parse_text(file_bytes: bytes) -> str:
    """解析 UTF-8 文字文件，保留 Markdown／純文字原貌。"""
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("文字文件必須使用 UTF-8 編碼。") from exc


def _terminate_process_tree(process: subprocess.Popen[bytes]) -> None:
    """連同 soffice fork 出來的子行程一起殺掉。

    只 kill 父行程的話，真正在轉檔的那個行程會變成孤兒繼續佔著 CPU；
    Windows 沒有 process group 的概念，退回單純 kill。
    """
    killpg = getattr(os, "killpg", None)
    getpgid = getattr(os, "getpgid", None)
    if killpg is not None and getpgid is not None:
        with contextlib.suppress(Exception):
            killpg(getpgid(process.pid), signal.SIGKILL)
            return
    with contextlib.suppress(Exception):
        process.kill()


def parse_doc(file_bytes: bytes) -> str:
    """透過隔離的 LibreOffice/soffice 將舊式 .doc 轉成純文字。"""
    from app.core.i18n import t

    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter:
        raise ValueError(t("rubric_parser.libreoffice_missing"))

    with tempfile.TemporaryDirectory(prefix="teacher-judge-doc-") as directory:
        input_path = Path(directory) / "source.doc"
        input_path.write_bytes(file_bytes)
        profile_dir = Path(directory) / "profile"
        profile_dir.mkdir()
        popen_kwargs: dict[str, Any] = {}
        if hasattr(os, "setsid"):
            # 自己一組 process group，逾時才殺得乾淨
            popen_kwargs["start_new_session"] = True
        process = subprocess.Popen(
            [
                converter,
                # 每次轉檔用獨立的使用者設定目錄：共用 profile 會讓並行轉檔
                # 互相搶鎖，也避免文件內容污染到下一次轉檔
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                directory,
                str(input_path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            **popen_kwargs,
        )
        try:
            process.communicate(timeout=DOC_CONVERT_TIMEOUT_SECONDS)
        except subprocess.TimeoutExpired:
            _terminate_process_tree(process)
            with contextlib.suppress(Exception):
                process.communicate(timeout=5)
            raise ValueError(t("rubric_parser.doc_conversion_timeout")) from None
        output_path = Path(directory) / "source.txt"
        if process.returncode != 0 or not output_path.exists():
            raise ValueError(t("rubric_parser.doc_conversion_failed"))
        return parse_text(output_path.read_bytes())


def parse_document(filename: str, file_bytes: bytes) -> str:
    """根據副檔名選擇解析器，回傳純文字（Markdown 格式）。"""
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        return parse_text(file_bytes)
    elif suffix == ".doc":
        return parse_doc(file_bytes)
    elif suffix == ".docx":
        return parse_docx(file_bytes)
    elif suffix == ".pdf":
        return parse_pdf(file_bytes)
    else:
        from app.core.i18n import t

        raise ValueError(t("rubric_parser.unsupported_format", suffix=suffix))


# ──────────────────────────────────────────────────────
# python-docx 工具：按文件順序迭代段落 + 表格
# ──────────────────────────────────────────────────────


def _iter_block_items(doc):  # type: ignore
    """
    按文件原始順序，逐一 yield 段落或表格。
    python-docx 的 doc.paragraphs 和 doc.tables 是分開的，
    此函式利用 XML element 順序確保正確排列。
    """

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            from docx.text.paragraph import Paragraph  # type: ignore

            yield {"type": "paragraph", "text": Paragraph(child, doc).text}
        elif tag == "tbl":
            from docx.table import Table  # type: ignore

            yield {"type": "table", "table": Table(child, doc)}
